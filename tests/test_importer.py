from datetime import date

from openpyxl import load_workbook

from app import importer
from app.models import ClassRecord, PurchaseRecord, Student


def _fill(path, rows, keep_example=False):
    importer.build_template(path)
    wb = load_workbook(path)
    ws = wb[importer.DATA_SHEET]
    keys = [key for key, *_rest in importer.COLUMNS]
    start = 3 if keep_example else 2
    if not keep_example:
        for col in range(1, len(keys) + 1):
            ws.cell(row=2, column=col).value = None
    for offset, row in enumerate(rows):
        for col, key in enumerate(keys, start=1):
            ws.cell(row=start + offset, column=col).value = row.get(key)
    wb.save(path)


def _row(**overrides):
    row = {
        "name": "陳大文",
        "phone_type": "手機",
        "phone": "0912345678",
        "gender": "男",
        "purchase_date": date(2026, 1, 1),
        "quantity": 10,
        "used": 3,
        "last_class_date": date(2026, 9, 1),
    }
    row.update(overrides)
    return row


def test_template_example_row_is_valid(app, tmp_path):
    path = tmp_path / "t.xlsx"
    importer.build_template(path)
    rows, errors = importer.read_rows(path)
    assert errors == []
    assert [d["name"] for _n, d in rows] == ["範例 王小明"]


def test_import_records_used_as_opening_deduction(app, db, tmp_path):
    path = tmp_path / "t.xlsx"
    _fill(path, [_row(price=12000), _row(name="林小美", phone_type=None, phone=None, used=0, quantity=5)])
    rows, errors = importer.read_rows(path)
    assert errors == []
    importer.import_rows(rows)

    chen = Student.query.filter_by(name="陳大文").one()
    assert (chen.total_purchased, chen.total_attended, chen.remaining) == (7, 0, 7)
    assert chen.gender == "male" and chen.phone_type == "mobile"
    (purchase,) = chen.purchase_records
    assert (purchase.quantity, purchase.opening_deduction, purchase.price) == (10, 3, 12000)
    lin = Student.query.filter_by(name="林小美").one()
    assert (lin.total_purchased, lin.remaining) == (5, 5)
    assert ClassRecord.query.count() == 0


def test_fully_used_row_imports_with_zero_remaining(app, db, tmp_path):
    path = tmp_path / "t.xlsx"
    _fill(path, [_row(quantity=10, used=10)])
    rows, errors = importer.read_rows(path)
    assert errors == []
    importer.import_rows(rows)
    student = Student.query.one()
    assert (student.total_purchased, student.remaining) == (0, 0)
    assert PurchaseRecord.query.one().quantity == 10


def test_validation_errors_are_reported_per_row(app, db, tmp_path):
    db.session.add(Student(name="舊學生", status="active"))
    db.session.commit()
    path = tmp_path / "t.xlsx"
    _fill(
        path,
        [
            _row(used=11),
            _row(name="舊學生"),
            _row(name="甲", quantity=None),
            _row(name="乙", phone="12345", phone_type="手機"),
            _row(name="丙", purchase_date="2026/13/40"),
            _row(name="丁"),
            _row(name="丁"),
        ],
    )
    rows, errors = importer.read_rows(path)
    assert [n for n, _m in errors] == [2, 3, 4, 5, 6, 8]
    assert "不能大於購買堂數" in errors[0][1]
    assert [d["name"] for _n, d in rows] == ["丁"]


def test_phone_losing_leading_zero_is_restored(app, tmp_path):
    path = tmp_path / "t.xlsx"
    _fill(path, [_row(phone=912345678)])
    rows, errors = importer.read_rows(path)
    assert errors == []
    assert rows[0][1]["phone"] == "0912345678"


def test_cli_dry_run_does_not_write(app, db, tmp_path):
    path = tmp_path / "t.xlsx"
    _fill(path, [_row()])
    runner = app.test_cli_runner()
    result = runner.invoke(args=["import-students", str(path)])
    assert result.exit_code == 0, result.output
    assert Student.query.count() == 0

    result = runner.invoke(args=["import-students", str(path), "--commit"])
    assert result.exit_code == 0, result.output
    assert Student.query.count() == 1


def test_cli_rejects_whole_batch_on_error(app, db, tmp_path):
    path = tmp_path / "t.xlsx"
    _fill(path, [_row(), _row(name="壞資料", used=99)])
    result = app.test_cli_runner().invoke(args=["import-students", str(path), "--commit"])
    assert result.exit_code != 0
    assert Student.query.count() == 0
