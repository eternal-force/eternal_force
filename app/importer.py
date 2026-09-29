"""學生堂數期初批量匯入：產生 Excel 範本、讀取並驗證、寫入資料庫。

一列代表一位學生：建立學生名冊與一筆購課紀錄(保留原始購買堂數與金額)。
「已使用堂數」記在購課紀錄的期初扣抵堂數，不產生上課紀錄，
因此總購買堂數＝剩餘堂數＝購買堂數－已使用堂數，已上堂數為 0。
"""

from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from . import services
from .extensions import db
from .models import PurchaseRecord, Student

DATA_SHEET = "學生匯入"
HELP_SHEET = "填寫說明"
MAX_ROWS = 500

# (欄位 key, 表頭, 必填, 欄寬, 說明)
COLUMNS = [
    ("name", "學生姓名", True, 14, "最多 20 字，不可含特殊符號；不可與系統內既有學生同名"),
    ("phone_type", "電話類型", False, 10, "手機 / 市話（有填電話時必填）"),
    ("phone", "聯絡電話", False, 15, "手機 0912345678；市話 02-12345678"),
    ("email", "Email", False, 24, "選填"),
    ("birthday", "生日", False, 12, "日期格式 YYYY-MM-DD"),
    ("gender", "性別", False, 8, "男 / 女 / 其他"),
    ("purchase_date", "購買日期", True, 12, "日期格式 YYYY-MM-DD"),
    ("quantity", "購買堂數", True, 10, "1–999 的整數"),
    ("price", "購買金額", False, 12, "選填，數字，最多 1,000,000"),
    ("used", "已使用堂數", True, 11, "0 到購買堂數之間的整數；沒用過填 0"),
    ("last_class_date", "最後上課日期", False, 13, "僅供參考，匯入時忽略"),
    ("notes", "備註", False, 24, "選填，寫入學生備註"),
]
REMAINING_HEADER = "剩餘堂數（自動計算，匯入時忽略）"

PHONE_TYPES = {"手機": "mobile", "市話": "landline"}
GENDERS = {"男": "male", "女": "female", "其他": "other"}

EXAMPLE_ROW = {
    "name": "範例 王小明",
    "phone_type": "手機",
    "phone": "0912345678",
    "email": "example@mail.com",
    "birthday": date(1995, 5, 20),
    "gender": "男",
    "purchase_date": date(2026, 1, 10),
    "quantity": 20,
    "price": 24000,
    "used": 8,
    "last_class_date": date(2026, 9, 20),
    "notes": "範例列，匯入前請刪除",
}


def build_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET

    font = Font(name="Arial", size=11)
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    required_fill = PatternFill("solid", fgColor="C0504D")
    optional_fill = PatternFill("solid", fgColor="4F81BD")
    calc_fill = PatternFill("solid", fgColor="808080")
    example_font = Font(name="Arial", size=11, italic=True, color="808080")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [(key, label + ("*" if required else ""), width, note) for key, label, required, width, note in COLUMNS]
    for col, (key, label, width, note) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = required_fill if label.endswith("*") else optional_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        cell.comment = Comment(note, "系統")
        ws.column_dimensions[cell.column_letter].width = width

    calc_col = len(COLUMNS) + 1
    calc_cell = ws.cell(row=1, column=calc_col, value=REMAINING_HEADER)
    calc_cell.font = header_font
    calc_cell.fill = calc_fill
    calc_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    calc_cell.border = border
    ws.column_dimensions[calc_cell.column_letter].width = 16
    ws.row_dimensions[1].height = 32

    col_of = {key: idx for idx, (key, *_rest) in enumerate(COLUMNS, start=1)}
    date_keys = {"birthday", "purchase_date", "last_class_date"}
    for idx, (key, *_rest) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=idx, value=EXAMPLE_ROW[key])
        cell.font = example_font
    for row in range(2, MAX_ROWS + 2):
        for idx, (key, *_rest) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row, column=idx)
            cell.border = border
            if row > 2:
                cell.font = font
            if key in date_keys:
                cell.number_format = "yyyy-mm-dd"
            elif key == "price":
                cell.number_format = "#,##0"
            elif key == "phone":
                cell.number_format = "@"
        q = ws.cell(row=row, column=col_of["quantity"]).coordinate
        u = ws.cell(row=row, column=col_of["used"]).coordinate
        calc = ws.cell(row=row, column=calc_col, value=f'=IF({q}="","",{q}-N({u}))')
        calc.font = example_font if row == 2 else font
        calc.border = border

    def add_list_validation(key, options):
        letter = ws.cell(row=1, column=col_of[key]).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.error = "請從下拉選單選擇"
        dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
        ws.add_data_validation(dv)

    def add_whole_validation(key, low, high, message):
        letter = ws.cell(row=1, column=col_of[key]).column_letter
        dv = DataValidation(type="whole", operator="between", formula1=str(low), formula2=str(high), allow_blank=True)
        dv.error = message
        dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
        ws.add_data_validation(dv)

    add_list_validation("phone_type", PHONE_TYPES)
    add_list_validation("gender", GENDERS)
    add_whole_validation("quantity", 1, 999, "購買堂數需為 1–999 的整數")
    add_whole_validation("used", 0, 999, "已使用堂數需為 0–999 的整數")

    ws.freeze_panes = "B2"

    help_ws = wb.create_sheet(HELP_SHEET)
    lines = [
        ("學生堂數期初批量匯入 — 填寫說明", True),
        ("", False),
        ("1. 在「學生匯入」工作表，一列填一位學生；第 2 列是範例，匯入前請刪除。", False),
        ("2. 紅色表頭（*）為必填，藍色為選填，灰色「剩餘堂數」自動計算僅供核對。", False),
        ("3. 日期請用 YYYY-MM-DD；電話類型與性別請用下拉選單。", False),
        ("4. 已使用堂數記為購課紀錄的「期初已扣」堂數，不建立上課紀錄；匯入後已上堂數為 0，總購買與剩餘堂數＝購買堂數－已使用堂數。", False),
        ("5. 學生姓名不可與系統內既有學生、或本檔其他列重複。", False),
        ("6. 匯入指令（先檢查、再正式寫入）：", False),
        ("   flask import-students 檔案.xlsx            → 只檢查，不寫入", False),
        ("   flask import-students 檔案.xlsx --commit   → 全部通過才寫入；任一列有誤則整批不寫入", False),
        ("", False),
        ("欄位", True),
    ]
    for key, label, required, _w, note in COLUMNS:
        lines.append((f"{label}{'（必填）' if required else ''}：{note}", False))
    for row, (text, bold) in enumerate(lines, start=1):
        cell = help_ws.cell(row=row, column=1, value=text)
        cell.font = Font(name="Arial", size=14 if row == 1 else 11, bold=bold)
    help_ws.column_dimensions["A"].width = 100

    wb.save(path)


# ---------- 讀取與驗證 ----------

def _text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip()
    return value or None


def _to_date(value, label, required=False):
    if value in (None, ""):
        if required:
            raise services.ValidationError(f"{label}為必填")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise services.ValidationError(f"{label}格式不正確，需為 YYYY-MM-DD")


def _to_int(value, label, low, high):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    try:
        number = int(str(value).strip())
    except (TypeError, ValueError):
        raise services.ValidationError(f"{label}必須是整數")
    if number < low or number > high:
        raise services.ValidationError(f"{label}需介於 {low} 到 {high}")
    return number


def _parse_row(raw):
    name = _text(raw.get("name"))
    if not name:
        raise services.ValidationError("學生姓名為必填")
    services.validate_student_name(name)

    phone = _text(raw.get("phone"))
    phone_type_label = _text(raw.get("phone_type"))
    if phone_type_label and phone_type_label not in PHONE_TYPES:
        raise services.ValidationError("電話類型需為「手機」或「市話」")
    phone_type = PHONE_TYPES.get(phone_type_label)
    if phone and len(phone) == 9 and phone.startswith("9"):
        phone = "0" + phone  # Excel 把 0912345678 當數字時會吃掉開頭的 0
    services.validate_phone_by_type(phone, phone_type)

    gender_label = _text(raw.get("gender"))
    if gender_label and gender_label not in GENDERS:
        raise services.ValidationError("性別需為「男」、「女」或「其他」")

    purchase_date = _to_date(raw.get("purchase_date"), "購買日期", required=True)
    if raw.get("quantity") in (None, ""):
        raise services.ValidationError("購買堂數為必填")
    quantity = _to_int(raw.get("quantity"), "購買堂數", 1, 999)
    if raw.get("used") in (None, ""):
        raise services.ValidationError("已使用堂數為必填（沒用過請填 0）")
    used = _to_int(raw.get("used"), "已使用堂數", 0, 999)
    if used > quantity:
        raise services.ValidationError("已使用堂數不能大於購買堂數")

    return {
        "name": name,
        "phone": phone,
        "phone_type": phone_type if phone else None,
        "email": _text(raw.get("email")),
        "birthday": _to_date(raw.get("birthday"), "生日"),
        "gender": GENDERS.get(gender_label),
        "purchase_date": purchase_date,
        "quantity": quantity,
        "price": services._validate_purchase_price(raw.get("price")),
        "used": used,
        "notes": _text(raw.get("notes")),
    }


def read_rows(path):
    """回傳 (rows, errors)。rows: [(excel 列號, 解析後資料)]；errors: [(excel 列號, 訊息)]。"""
    wb = load_workbook(path, data_only=True)
    if DATA_SHEET not in wb.sheetnames:
        return [], [(0, f"找不到工作表「{DATA_SHEET}」，請使用系統產生的範本")]
    ws = wb[DATA_SHEET]

    header_to_key = {label: key for key, label, *_rest in COLUMNS}
    positions = {}
    for idx, cell in enumerate(ws[1]):
        label = (_text(cell.value) or "").rstrip("*")
        if label in header_to_key:
            positions[header_to_key[label]] = idx
    missing = [label for key, label, *_rest in COLUMNS if key not in positions]
    if missing:
        return [], [(1, "表頭缺少欄位：" + "、".join(missing))]

    rows, errors = [], []
    existing = {name for (name,) in db.session.query(Student.name).all()}
    seen = {}
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw = {key: values[idx] if idx < len(values) else None for key, idx in positions.items()}
        if all(_text(v) is None for v in raw.values()):
            continue
        try:
            data = _parse_row(raw)
        except services.ValidationError as exc:
            errors.append((row_no, exc.message))
            continue
        if data["name"] in existing:
            errors.append((row_no, f"系統內已有學生「{data['name']}」"))
            continue
        if data["name"] in seen:
            errors.append((row_no, f"與第 {seen[data['name']]} 列學生姓名重複"))
            continue
        seen[data["name"]] = row_no
        rows.append((row_no, data))
    return rows, errors


def import_rows(rows):
    """在同一個交易內寫入全部資料；任何例外都整批回滾。"""
    try:
        for _row_no, data in rows:
            student = Student(
                name=data["name"],
                phone=data["phone"],
                phone_type=data["phone_type"],
                email=data["email"],
                birthday=data["birthday"],
                gender=data["gender"],
                enrollment_date=data["purchase_date"],
                status="active",
                notes=data["notes"],
            )
            student.purchase_records.append(
                PurchaseRecord(
                    purchase_date=data["purchase_date"],
                    quantity=data["quantity"],
                    opening_deduction=data["used"],
                    price=data["price"],
                    notes="期初匯入",
                )
            )
            db.session.add(student)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
